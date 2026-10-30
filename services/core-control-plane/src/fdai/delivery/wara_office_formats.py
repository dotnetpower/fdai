"""Delivery-owned Office encoders registered only by the WARA file sink."""

from __future__ import annotations

import csv
import io
import json
import textwrap
from typing import Literal

from fdai_service_contracts.wara_guidance import GuidanceDiagram, ReferenceImage
from openpyxl import Workbook
from pptx import Presentation
from pptx.presentation import Presentation as PresentationDocument
from pptx.util import Inches, Pt

from fdai.core.reporting.formats.csv_format import CsvFormatEncoder
from fdai.core.reporting.formats.json_format import JsonFormatEncoder
from fdai.core.reporting.models import RenderedReport
from fdai.delivery.wara_guidance_presentation import (
    guidance_diagram_slides,
    guidance_reference_image_slides,
    guidance_text_slides,
)

_MAXIMUM_SLIDES = 1_000


def _admit_size(report: RenderedReport) -> None:
    if len(JsonFormatEncoder().encode(report)) > 8_388_608:
        raise ValueError("Office report exceeds the encoded byte limit")


def _metadata(report: RenderedReport) -> dict[str, object]:
    value = report.to_dict()
    value.pop("widgets")
    return value


class XlsxFormatEncoder:
    """Export tabular data and complete report metadata as text-only cells."""

    name = "xlsx"
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def encode(self, report: RenderedReport) -> bytes:
        _admit_size(report)
        workbook = Workbook()
        sheet = workbook.active
        if sheet is None:
            raise RuntimeError("report workbook has no worksheet")
        sheet.title = "Report"
        rows = csv.reader(io.StringIO(CsvFormatEncoder().encode(report).decode("utf-8")))
        for row in rows:
            if any(len(value) > 32_767 for value in row):
                raise ValueError("report cell exceeds the Excel text limit")
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        metadata = workbook.create_sheet("Metadata")
        for key, value in _metadata(report).items():
            text = json.dumps(value, ensure_ascii=False, sort_keys=True)
            for offset in range(0, len(text), 32_000):
                metadata.append([key, offset // 32_000, text[offset : offset + 32_000]])
                metadata.cell(metadata.max_row, 1).data_type = "s"
                metadata.cell(metadata.max_row, 3).data_type = "s"
        status = workbook.create_sheet("Widget status")
        status.append(["widget_id", "error"])
        pending = list(report.widgets)
        while pending:
            widget = pending.pop(0)
            if len(widget.id) > 32_767 or len(widget.error or "") > 32_767:
                raise ValueError("report cell exceeds the Excel text limit")
            status.append([widget.id, widget.error or ""])
            for cell in status[status.max_row]:
                cell.data_type = "s"
            pending[0:0] = widget.children
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()


class PptxFormatEncoder:
    """Paginate complete report content instead of silently shortening slides."""

    name = "pptx"
    content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"

    def encode(self, report: RenderedReport) -> bytes:
        _admit_size(report)
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        if "public-guidance" in report.tags:
            return self._guidance(presentation, report)
        metadata = json.dumps(_metadata(report), ensure_ascii=False, sort_keys=True, indent=2)
        self._slides(presentation, report.name, metadata)
        pending = list(report.widgets)
        while pending:
            widget = pending.pop(0)
            material = {
                "widget_id": widget.id,
                "type": widget.type,
                "error": widget.error,
                "data": widget.data,
            }
            self._slides(
                presentation,
                widget.title,
                json.dumps(material, ensure_ascii=False, sort_keys=True, indent=2),
            )
            pending[0:0] = widget.children
        output = io.BytesIO()
        presentation.save(output)
        return output.getvalue()

    @staticmethod
    def _guidance(presentation: PresentationDocument, report: RenderedReport) -> bytes:
        selected_locale = report.variables.get("locale", "en")
        if selected_locale not in {"en", "ko"}:
            raise ValueError("WARA guidance presentation locale is invalid")
        locale: Literal["en", "ko"] = "ko" if selected_locale == "ko" else "en"
        guidance_text_slides(
            presentation,
            report.name,
            report.description
            + "\n\n"
            + "\n".join(f"{key}: {value}" for key, value in report.variables.items()),
        )
        for widget in report.widgets:
            if widget.error is not None:
                raise ValueError("WARA guidance presentation cannot hide a widget error")
            if widget.type == "wara_guidance_reference_image":
                image = ReferenceImage.model_validate(widget.data.get("reference_image"))
                try:
                    guidance_reference_image_slides(presentation, image, locale)
                except (OSError, ValueError) as exc:
                    raise ValueError(f"WARA reference image export failed: {exc}") from exc
            elif widget.type == "wara_guidance_diagram":
                diagram = GuidanceDiagram.model_validate(widget.data.get("diagram"))
                guidance_diagram_slides(presentation, diagram, locale)
            elif widget.type == "wara_guidance_text":
                body = widget.data.get("body")
                if not isinstance(body, str):
                    raise ValueError("WARA guidance presentation text is malformed")
                guidance_text_slides(presentation, widget.title, body)
            else:
                raise ValueError("WARA guidance presentation widget is unsupported")
        output = io.BytesIO()
        presentation.save(output)
        return output.getvalue()

    @staticmethod
    def _slides(presentation: PresentationDocument, title: str, text: str) -> None:
        lines = [
            wrapped
            for line in text.splitlines()
            for wrapped in (textwrap.wrap(line, width=100, replace_whitespace=False) or [""])
        ]
        chunks = [lines[offset : offset + 22] for offset in range(0, len(lines), 22)]
        if len(presentation.slides) + len(chunks) > _MAXIMUM_SLIDES:
            raise ValueError("PowerPoint report exceeds the slide limit")
        for index, chunk in enumerate(chunks, start=1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            heading = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.7))
            heading.text_frame.text = f"{title} ({index}/{len(chunks)})"
            heading.text_frame.paragraphs[0].font.size = Pt(24)
            body = slide.shapes.add_textbox(Inches(0.5), Inches(1.0), Inches(12.3), Inches(6.1))
            body.text_frame.word_wrap = True
            for ordinal, line in enumerate(chunk):
                paragraph = (
                    body.text_frame.paragraphs[0]
                    if ordinal == 0
                    else body.text_frame.add_paragraph()
                )
                paragraph.text = line
                paragraph.font.size = Pt(14)


__all__ = ["PptxFormatEncoder", "XlsxFormatEncoder"]
