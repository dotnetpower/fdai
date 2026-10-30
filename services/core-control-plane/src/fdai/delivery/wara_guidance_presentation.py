"""Original diagram and readable prose slides for public WARA guidance."""

from __future__ import annotations

import textwrap
import unicodedata
from io import BytesIO
from typing import Literal

from fdai_service_contracts.reference_image_assets import load_packaged_reference_image
from fdai_service_contracts.wara_guidance import GuidanceDiagram, GuidanceText, ReferenceImage
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.presentation import Presentation
from pptx.slide import Slide
from pptx.util import Inches, Pt

from fdai.core.wara.guidance_report import guidance_text

_SLIDE_LIMIT = 1_000
_INK = MSO_THEME_COLOR.DARK_1


def _display_columns(text: str) -> int:
    return sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in text)


def guidance_text_pages(text: str) -> tuple[tuple[str, ...], ...]:
    """Reserve visual rows for full-width Korean glyphs and unbroken source URLs."""

    lines = [
        wrapped
        for line in text.splitlines()
        for wrapped in (
            textwrap.wrap(
                line,
                width=88,
                replace_whitespace=False,
                break_long_words=False,
                break_on_hyphens=False,
            )
            or [""]
        )
    ]
    pages: list[tuple[str, ...]] = []
    current: list[str] = []
    rows = 0
    for line in lines:
        required_rows = max(1, (_display_columns(line) + 87) // 88)
        if required_rows > 13:
            raise ValueError("WARA guidance paragraph exceeds its readable line budget")
        if rows + required_rows > 13:
            pages.append(tuple(current))
            current, rows = [], 0
        current.append(line)
        rows += required_rows
    if current or not pages:
        pages.append(tuple(current))
    return tuple(pages)


def guidance_diagram_geometry(
    diagram: GuidanceDiagram,
) -> tuple[dict[str, tuple[float, float]], tuple[tuple[tuple[float, float], ...], ...]]:
    """Keep routes in row gutters so no connector traverses an unrelated node."""

    width, height = 3.65, 1.15
    positions = {
        node.id: (0.55 + (index % 3) * 4.18, 1.65 + (index // 3) * 1.72)
        for index, node in enumerate(diagram.nodes)
    }
    routes = []
    for index, edge in enumerate(diagram.edges):
        source_x, source_y = positions[edge.from_]
        target_x, target_y = positions[edge.to]
        points: tuple[tuple[float, float], ...]
        if source_y == target_y and abs(source_x - target_x) < 4.2:
            forward = target_x > source_x
            points = (
                (source_x + width if forward else source_x, source_y + height / 2),
                (target_x if forward else target_x + width, target_y + height / 2),
            )
        else:
            downward = target_y > source_y
            source_port = (
                source_x + width / 2,
                source_y + height if downward else source_y,
            )
            target_port = (
                target_x + width / 2,
                target_y if downward else target_y + height if source_y != target_y else target_y,
            )
            source_gap = source_port[1] + (0.25 if downward else -0.25)
            target_gap = target_port[1] + (-0.25 if downward or source_y == target_y else 0.25)
            gutter_x = 12.8 + (index % 3) * 0.1
            points = (
                source_port,
                (source_port[0], source_gap),
                (gutter_x, source_gap),
                (gutter_x, target_gap),
                (target_port[0], target_gap),
                target_port,
            )
        compact = [points[0]]
        for point in points[1:]:
            if point != compact[-1]:
                compact.append(point)
        routes.append(tuple(compact))
    return positions, tuple(routes)


def _textbox(
    slide: Slide, text: str, *, x: float, y: float, width: float, height: float, size: int
) -> None:
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.text_frame.word_wrap = True
    shape.text_frame.margin_left = Inches(0.06)
    shape.text_frame.margin_right = Inches(0.06)
    shape.text_frame.margin_top = Inches(0.01)
    shape.text_frame.margin_bottom = Inches(0.01)
    for ordinal, line in enumerate(text.splitlines() or [""]):
        paragraph = (
            shape.text_frame.paragraphs[0] if ordinal == 0 else shape.text_frame.add_paragraph()
        )
        paragraph.text = line
        paragraph.font.size = Pt(size)
        paragraph.font.color.theme_color = _INK
        paragraph.space_after = Pt(0)
        paragraph.space_before = Pt(0)
        paragraph.line_spacing = 1.1


def guidance_text_slides(presentation: Presentation, title: str, text: str) -> None:
    """Use a primary-body floor of 18 points and page rather than shrink text."""

    title_lines = textwrap.wrap(title, width=70, replace_whitespace=False) or [""]
    if len(title_lines) > 3:
        raise ValueError("WARA guidance presentation title exceeds its line budget")
    chunks = guidance_text_pages(text)
    if len(presentation.slides) + len(chunks) > _SLIDE_LIMIT:
        raise ValueError("WARA guidance presentation exceeds its slide limit")
    for number, chunk in enumerate(chunks, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _textbox(slide, "\n".join(title_lines), x=0.5, y=0.25, width=12.3, height=1.35, size=28)
        _textbox(slide, "\n".join(chunk), x=0.55, y=1.75, width=12.2, height=5.1, size=18)
        _textbox(
            slide,
            f"Public guidance | {number}/{len(chunks)} | No execution authority",
            x=0.55,
            y=7.0,
            width=12.2,
            height=0.25,
            size=10,
        )


def guidance_diagram_slides(
    presentation: Presentation, diagram: GuidanceDiagram, locale: Literal["en", "ko"]
) -> None:
    """Draw native vector nodes and connected arrows, plus complete accessible explanations."""

    if len(presentation.slides) >= _SLIDE_LIMIT:
        raise ValueError("WARA guidance presentation exceeds its slide limit")
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(
        slide,
        guidance_text(diagram.title, locale),
        x=0.5,
        y=0.2,
        width=12.3,
        height=1.0,
        size=26,
    )
    columns = 3
    rows = (len(diagram.nodes) + columns - 1) // columns
    width, height = 3.65, 1.15
    positions, routes = guidance_diagram_geometry(diagram)
    for edge, points in zip(diagram.edges, routes, strict=True):
        for segment, (start, end) in enumerate(zip(points[:-1], points[1:], strict=True)):
            connector = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(start[0]),
                Inches(start[1]),
                Inches(end[0]),
                Inches(end[1]),
            )
            connector.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
            connector.line.width = Pt(1.5)
            if edge.style == "dashed":
                connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            if segment == len(points) - 2:
                arrow = OxmlElement("a:tailEnd")
                arrow.set("type", "triangle")
                connector.line._get_or_add_ln().append(arrow)
    for node in diagram.nodes:
        x, y = positions[node.id]
        shape = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height),
        )
        shape.fill.solid()
        shape.fill.fore_color.theme_color = MSO_THEME_COLOR.LIGHT_1
        shape.line.color.theme_color = MSO_THEME_COLOR.ACCENT_1
        shape.text_frame.word_wrap = True
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        shape.text_frame.margin_left = Inches(0.12)
        shape.text_frame.margin_right = Inches(0.12)
        shape.text_frame.margin_top = Inches(0.01)
        shape.text_frame.margin_bottom = Inches(0.01)
        label = guidance_text(node.label, locale)
        if _display_columns(label) > 108:
            raise ValueError("WARA guidance diagram label exceeds its readable text budget")
        shape.text_frame.text = label
        for paragraph in shape.text_frame.paragraphs:
            paragraph.font.size = Pt(18)
            paragraph.font.color.theme_color = _INK
            paragraph.space_after = Pt(0)
            paragraph.space_before = Pt(0)
            paragraph.line_spacing = 1.05
    _textbox(
        slide,
        "구성 예시: 실제 관찰·자동 장애 조치·실행 권한이 아닙니다."
        if locale == "ko"
        else (
            "Illustrative configuration: not an observation, "
            "automatic failover or execution authority."
        ),
        x=0.55,
        y=1.65 + rows * 1.72,
        width=12.2,
        height=0.6,
        size=14,
    )
    labels = {node.id: guidance_text(node.label, locale) for node in diagram.nodes}
    details = [
        guidance_text(diagram.description, locale),
        *[
            f"{guidance_text(node.label, locale)}: {guidance_text(node.detail, locale)}"
            for node in diagram.nodes
        ],
        *[
            f"{labels[edge.from_]} -> {labels[edge.to]}"
            f": {guidance_text(edge.label, locale)} ({edge.style})"
            for edge in diagram.edges
        ],
    ]
    guidance_text_slides(
        presentation,
        "도표 설명 및 연결선" if locale == "ko" else "Diagram explanation and connections",
        "\n\n".join(details),
    )


def guidance_reference_image_slides(
    presentation: Presentation,
    image: ReferenceImage,
    locale: Literal["en", "ko"],
) -> None:
    """Embed a verified package asset before the supplementary authored diagram."""

    if len(presentation.slides) >= _SLIDE_LIMIT:
        raise ValueError("WARA guidance presentation exceeds its slide limit")
    asset = load_packaged_reference_image(image)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _textbox(
        slide,
        "공식 참고 이미지" if locale == "ko" else "Official reference image",
        x=0.5,
        y=0.3,
        width=12.3,
        height=0.7,
        size=26,
    )
    max_width, max_height = 11.8, 5.0
    scale = min(max_width / asset.width, max_height / asset.height)
    width, height = asset.width * scale, asset.height * scale
    picture = slide.shapes.add_picture(
        BytesIO(asset.content),
        Inches((13.333 - width) / 2),
        Inches(1.15 + (max_height - height) / 2),
        width=Inches(width),
        height=Inches(height),
    )
    picture._element.nvPicPr.cNvPr.set("descr", _localized_image_text(image.alt, locale))
    _textbox(
        slide,
        f"{image.attribution} | {image.license} | {image.source_document_url}",
        x=0.5,
        y=6.5,
        width=12.3,
        height=0.7,
        size=12,
    )
    provenance = "\n".join(
        (
            f"{'설명' if locale == 'ko' else 'Caption'}: "
            f"{_localized_image_text(image.caption, locale)}",
            f"{'대체 텍스트' if locale == 'ko' else 'Alt text'}: "
            f"{_localized_image_text(image.alt, locale)}",
            f"{'이미지 내 텍스트' if locale == 'ko' else 'Text in image'}: "
            f"{_localized_image_text(image.english_in_image_notice, locale)}",
            f"{'저작자 표시' if locale == 'ko' else 'Attribution'}: {image.attribution}",
            f"{'라이선스' if locale == 'ko' else 'License'}: {image.license} ({image.license_url})",
            f"{'문서 출처' if locale == 'ko' else 'Source document'}: {image.source_document_url}",
            f"{'이미지 출처' if locale == 'ko' else 'Image source'}: {image.image_url}",
            f"{'원본 개정' if locale == 'ko' else 'Source revision'}: {image.source_revision}",
            f"{'자산 다이제스트' if locale == 'ko' else 'Asset digest'}: {image.digest}",
        )
    )
    guidance_text_slides(
        presentation,
        "참고 이미지 출처 및 설명" if locale == "ko" else "Reference image attribution",
        provenance,
    )


def _localized_image_text(image_text: GuidanceText, locale: Literal["en", "ko"]) -> str:
    return image_text.ko if locale == "ko" and image_text.ko is not None else image_text.en


__all__ = [
    "guidance_diagram_geometry",
    "guidance_diagram_slides",
    "guidance_reference_image_slides",
    "guidance_text_pages",
    "guidance_text_slides",
]
