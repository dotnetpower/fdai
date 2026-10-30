"""Public guide exports contain original diagrams, readable text and no installation facts."""

from __future__ import annotations

import hashlib
import json
import struct
import zlib
from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

import pytest
from fdai.core.wara.guidance_report import build_wara_guidance_report
from fdai.delivery.wara_guidance_presentation import (
    guidance_diagram_geometry,
    guidance_reference_image_slides,
    guidance_text_pages,
)
from fdai.delivery.wara_office_formats import PptxFormatEncoder
from fdai_service_contracts.catalog_canonical import canonical_digest
from fdai_service_contracts.reference_image_assets import resources as image_resources
from fdai_service_contracts.reference_image_assets import verify_reference_image_bytes
from fdai_service_contracts.wara_guidance import (
    ReferenceImage,
    WaraGuidance,
    packaged_wara_guidance,
)
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _png_chunk(kind: bytes, data: bytes) -> bytes:
    return (
        struct.pack(">I", len(data))
        + kind
        + data
        + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    )


def _test_png() -> bytes:
    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\x00\x00\x00")
    return (
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", header)
        + _png_chunk(b"IDAT", pixels)
        + _png_chunk(b"IEND", b"")
    )


def _reference_image(content: bytes) -> ReferenceImage:
    digest = hashlib.sha256(content).hexdigest()
    return ReferenceImage.model_validate(
        {
            "asset_id": f"sha256-{digest}",
            "digest": f"sha256:{digest}",
            "source_document_url": (
                "https://learn.microsoft.com/en-us/azure/backup/backup-azure-vms-introduction"
            ),
            "image_url": (
                "https://learn.microsoft.com/en-us/azure/backup/media/"
                "backup-azure-vms-introduction/vmbackup-architecture.png"
            ),
            "source_revision": "a" * 40,
            "attribution": "Microsoft, About Azure VM backup - Azure Backup",
            "license": "CC-BY-4.0",
            "license_url": "https://creativecommons.org/licenses/by/4.0/",
            "alt": {"en": "A reference image.", "ko": "참고 이미지."},
            "caption": {"en": "Backup flow.", "ko": "백업 흐름."},
            "english_in_image_notice": {
                "en": "Text in the image is in English.",
                "ko": "이미지 내 텍스트는 영어입니다.",
            },
        }
    )


@pytest.fixture
def synthetic_packaged_image(monkeypatch):
    root = Path(f".fdai-reference-image-test-{uuid4().hex}")
    directory = root / "data/wara-guidance-assets"
    directory.mkdir(parents=True)
    content = _test_png()
    image = _reference_image(content)
    asset = directory / f"{image.asset_id}.png"
    asset.write_bytes(content)
    original_files = image_resources.files

    class RedirectedResources:
        def joinpath(self, *parts: str):
            if parts == ("data", "wara-guidance-assets"):
                return directory
            return original_files("fdai_service_contracts").joinpath(*parts)

    monkeypatch.setattr(image_resources, "files", lambda _: RedirectedResources())
    try:
        yield content, image, asset
    finally:
        asset.unlink(missing_ok=True)
        directory.rmdir()
        directory.parent.rmdir()
        root.rmdir()


def _guide_with_reference_image(image: ReferenceImage) -> WaraGuidance:
    raw = (
        packaged_wara_guidance()
        .entries[0]
        .model_dump(mode="json", by_alias=True, exclude_none=True)
    )
    raw["reference_image"] = image.model_dump(mode="json", by_alias=True, exclude_none=True)
    material = {key: value for key, value in raw.items() if key != "content_digest"}
    raw["content_digest"] = "sha256:" + canonical_digest(material)
    return WaraGuidance.model_validate(raw)


@pytest.mark.parametrize(
    "needle",
    [
        "Use virtual network links to link global Private DNS Zones",
        "Use maintenance configurations for the Dedicated",
        "Deploy AOAI Service in multiple regions",
    ],
)
@pytest.mark.parametrize("locale", ["en", "ko"])
def test_benchmark_export_draws_real_vector_diagrams_with_complete_text(needle, locale) -> None:
    catalog = packaged_wara_guidance()
    framework = json.loads(
        (
            Path(__file__).resolve().parents[4] / "rule-catalog/collected/wara-aprl/azure-wara.json"
        ).read_text()
    )
    identity = next(
        control["id"]
        for area in framework["areas"]
        for control in area["controls"]
        if control["title"].startswith(needle)
    )
    guide = catalog.resolve(identity)
    assert guide is not None
    report = build_wara_guidance_report(
        guide, locale=locale, generated_at=datetime(2026, 10, 30, tzinfo=UTC)
    )
    assert guide.reference_image is None
    assert not any(widget.type == "wara_guidance_reference_image" for widget in report.widgets)
    deck = Presentation(BytesIO(PptxFormatEncoder().encode(report)))
    assert not any(
        shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        for slide in deck.slides
        for shape in slide.shapes
    )
    assert (
        deck.slide_width * 9 == deck.slide_height * 16
        or abs(deck.slide_width / deck.slide_height - 16 / 9) < 0.001
    )
    nodes = [
        shape
        for slide in deck.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
    ]
    connectors = [
        shape
        for slide in deck.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.LINE
    ]
    assert len(nodes) == len(guide.diagram.nodes)
    assert len(connectors) >= len(guide.diagram.edges)
    assert sum(
        bool(connector._element.xpath(".//a:tailEnd[@type='triangle']")) for connector in connectors
    ) == len(guide.diagram.edges)
    assert len(deck.slides) >= 9
    text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert guide.content_digest in text
    assert guide.recommendation_id in text
    assert all(source.url in text for source in guide.sources)
    assert all(
        (node.label.ko if locale == "ko" and node.label.ko else node.label.en) in text
        for node in guide.diagram.nodes
    )
    assert "scope_digest" not in report.variables
    assert "satisfaction" not in report.variables
    assert report.variables["execution_authority"] == "false"
    for slide in deck.slides:
        for shape in slide.shapes:
            assert 0 <= shape.left < deck.slide_width
            assert 0 <= shape.top < deck.slide_height
            assert shape.left + shape.width <= deck.slide_width
            assert shape.top + shape.height <= deck.slide_height


def test_public_guidance_does_not_replace_disabled_history_with_current_advice() -> None:
    guide = next(
        entry for entry in packaged_wara_guidance().entries if entry.lifecycle == "disabled"
    )
    report = build_wara_guidance_report(
        guide, locale="en", generated_at=datetime(2026, 10, 30, tzinfo=UTC)
    )
    assert "disabled" in report.tags
    assert report.variables["lifecycle"] == "disabled"
    assert report.variables["execution_authority"] == "false"


def test_public_export_does_not_hide_unavailable_widgets() -> None:
    guide = packaged_wara_guidance().entries[0]
    report = build_wara_guidance_report(
        guide, locale="en", generated_at=datetime(2026, 10, 30, tzinfo=UTC)
    )
    failed = replace(report.widgets[0], error="source_unavailable")
    with pytest.raises(ValueError, match="widget error"):
        PptxFormatEncoder().encode(replace(report, widgets=(failed,)))


def test_reference_image_provenance_is_structured_in_json_and_paginated_in_pptx(
    synthetic_packaged_image,
) -> None:
    _, image, _ = synthetic_packaged_image
    guide_with_image = _guide_with_reference_image(image)
    report = build_wara_guidance_report(
        guide_with_image, locale="ko", generated_at=datetime(2026, 10, 30, tzinfo=UTC)
    )
    sources = next(widget for widget in report.widgets if widget.id == "sources")
    expected_image = image.model_dump(mode="json", by_alias=True, exclude_none=True)
    assert sources.data["reference_image"] == expected_image
    assert report.to_dict()["widgets"][-1]["data"]["reference_image"] == expected_image
    text = sources.data["body"]
    for expected in (
        "참고 이미지.",
        "백업 흐름.",
        "이미지 내 텍스트는 영어입니다.",
        image.attribution,
        image.source_document_url,
        image.image_url,
        image.license_url,
        image.digest,
    ):
        assert expected in text
    deck = Presentation(BytesIO(PptxFormatEncoder().encode(report)))
    exported_text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert image.attribution in exported_text
    assert image.image_url in exported_text
    assert "참고 이미지." in exported_text


@pytest.mark.parametrize("locale", ["en", "ko"])
def test_actual_guidance_pptx_embeds_packaged_image_and_complete_provenance(
    synthetic_packaged_image,
    locale,
) -> None:
    content, image, _ = synthetic_packaged_image
    report = build_wara_guidance_report(
        _guide_with_reference_image(image),
        locale=locale,
        generated_at=datetime(2026, 10, 30, tzinfo=UTC),
    )
    output = PptxFormatEncoder().encode(report)
    deck = Presentation(BytesIO(output))
    pictures = [
        shape
        for slide in deck.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1
    assert pictures[0].image.blob == content
    localized_alt = image.alt.ko if locale == "ko" else image.alt.en
    assert pictures[0]._element.nvPicPr.cNvPr.get("descr") == localized_alt
    text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    for expected in (
        image.caption.ko if locale == "ko" else image.caption.en,
        localized_alt,
        image.english_in_image_notice.ko if locale == "ko" else image.english_in_image_notice.en,
        image.source_document_url,
        image.image_url,
        image.source_revision,
        image.attribution,
        image.license,
        image.license_url,
        image.digest,
    ):
        assert expected in text
    with ZipFile(BytesIO(output)) as package:
        media = [name for name in package.namelist() if name.startswith("ppt/media/")]
        assert len(media) == 1
        assert package.read(media[0]) == content
        relationships = [
            package.read(name).decode()
            for name in package.namelist()
            if name.startswith("ppt/slides/_rels/") and name.endswith(".rels")
        ]
        assert any(
            '/image"' in relationship and "../media/" in relationship
            for relationship in relationships
        )


@pytest.mark.parametrize("failure", ["missing", "tampered"])
def test_actual_guidance_pptx_rejects_unavailable_or_invalid_reference_image(
    synthetic_packaged_image,
    failure,
) -> None:
    _, image, asset = synthetic_packaged_image
    report = build_wara_guidance_report(
        _guide_with_reference_image(image),
        locale="en",
        generated_at=datetime(2026, 10, 30, tzinfo=UTC),
    )
    if failure == "missing":
        asset.unlink()
    else:
        asset.write_bytes(b"not the approved PNG")
    with pytest.raises(ValueError, match="WARA reference image export failed"):
        PptxFormatEncoder().encode(report)


def test_reference_image_presentation_helper_embeds_verified_local_bytes(monkeypatch) -> None:
    content = _test_png()
    image = _reference_image(content)
    monkeypatch.setattr(
        "fdai.delivery.wara_guidance_presentation.load_packaged_reference_image",
        lambda value: verify_reference_image_bytes(value, content),
    )
    deck = Presentation()
    guidance_reference_image_slides(deck, image, "ko")
    pictures = [
        shape
        for slide in deck.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert len(pictures) == 1
    assert len(deck.slides) >= 2
    text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert image.attribution in text
    assert image.source_document_url in text
    assert image.image_url in text
    assert image.license_url in text
    assert image.source_revision in text
    assert "참고 이미지." in text
    assert "이미지 내 텍스트는 영어입니다." in text


def test_all_456_presentation_diagrams_keep_connectors_out_of_other_nodes() -> None:
    for guide in packaged_wara_guidance().entries:
        positions, routes = guidance_diagram_geometry(guide.diagram)
        assert len(routes) == len(guide.diagram.edges)
        for edge, points in zip(guide.diagram.edges, routes, strict=True):
            for start, end in zip(points[:-1], points[1:], strict=True):
                assert start[0] == end[0] or start[1] == end[1]
                for identity, (left, top) in positions.items():
                    if identity in {edge.from_, edge.to}:
                        continue
                    right, bottom = left + 3.65, top + 1.15
                    if start[1] == end[1]:
                        enters = (
                            top < start[1] < bottom
                            and min(start[0], end[0]) < right
                            and max(start[0], end[0]) > left
                        )
                    else:
                        enters = (
                            left < start[0] < right
                            and min(start[1], end[1]) < bottom
                            and max(start[1], end[1]) > top
                        )
                    assert not enters, (guide.recommendation_id, identity, edge)


def test_korean_prose_reserves_full_width_glyph_rows_instead_of_clipping() -> None:
    line = "검토한 구성과 현재 관찰을 구분하고 명시적인 조건을 확인합니다. " * 2
    text = "\n".join([line] * 13)
    pages = guidance_text_pages(text)
    assert len(pages) >= 2
    assert "".join("".join(page) for page in pages).replace(" ", "") == text.replace(
        "\n", ""
    ).replace(" ", "")
