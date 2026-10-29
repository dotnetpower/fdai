"""Office exports retain complete data and do not create spreadsheet formulas."""

from __future__ import annotations

from dataclasses import replace
from datetime import UTC, datetime
from io import BytesIO

import pytest
from fdai.core.reporting.models import RenderedReport, RenderedWidget
from fdai.delivery.wara_office_formats import PptxFormatEncoder, XlsxFormatEncoder
from openpyxl import load_workbook
from pptx import Presentation


def _report() -> RenderedReport:
    at = datetime(2026, 10, 30, tzinfo=UTC)
    return RenderedReport(
        id="wara",
        version="1.0.0",
        name="WARA",
        description="No execution authority.",
        generated_at=at,
        time_range=(at, at),
        variables={"execution_authority": "false"},
        widgets=(
            RenderedWidget(
                id="controls",
                type="table",
                title="Every control",
                data={
                    "columns": ["id", "status", "text"],
                    "rows": [
                        {"id": f"marker-{i:03d}", "status": "unknown", "text": "=1+1"}
                        for i in range(393)
                    ],
                },
            ),
            RenderedWidget(
                id="error",
                type="table",
                title="Unavailable source",
                data={},
                error="source_unavailable",
            ),
        ),
    )


def test_xlsx_roundtrip_retains_all_controls_and_error_state() -> None:
    book = load_workbook(BytesIO(XlsxFormatEncoder().encode(_report())))
    rows = list(book["Report"].values)
    assert len([row for row in rows if row[0] == "controls"]) == 393
    assert rows[-2][3] == "marker-392"
    assert all(cell.data_type != "f" for row in book["Report"] for cell in row)
    assert ("error", "source_unavailable") in list(book["Widget status"].values)
    assert "execution_authority" in str(list(book["Metadata"].values))


def test_pptx_roundtrip_paginates_without_losing_control_or_error_text() -> None:
    slides = Presentation(BytesIO(PptxFormatEncoder().encode(_report())))
    text = "\n".join(
        shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert len(slides.slides) > 1
    assert all(f"marker-{i:03d}" in text for i in range(393))
    assert "source_unavailable" in text
    assert "execution_authority" in text


def test_oversized_excel_cell_is_an_explicit_error_not_truncated() -> None:
    report = _report()
    widget = replace(
        report.widgets[0],
        data={"columns": ["value"], "rows": [{"value": "x" * 32_768}]},
    )
    with pytest.raises(ValueError, match="Excel text limit"):
        XlsxFormatEncoder().encode(replace(report, widgets=(widget,)))


@pytest.mark.parametrize("encoder", [XlsxFormatEncoder(), PptxFormatEncoder()])
def test_office_reports_have_an_explicit_eight_mib_input_ceiling(encoder) -> None:
    report = _report()
    oversized = replace(report, description="x" * 8_388_609)
    with pytest.raises(ValueError, match="encoded byte limit"):
        encoder.encode(oversized)


def test_powerpoint_slide_budget_fails_instead_of_omitting_content() -> None:
    report = replace(_report(), description="x" * 2_200_001, widgets=())
    with pytest.raises(ValueError, match="slide limit"):
        PptxFormatEncoder().encode(report)
