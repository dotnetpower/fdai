"""WARA exports preserve every recommendation and never trust imported context as a result."""

from __future__ import annotations

import csv
import json
import stat
from dataclasses import replace
from datetime import UTC, datetime
from pathlib import Path

import pytest
from fdai.core.wara import WaraAssessmentRequest, WaraAssessmentRuntime, WaraScopedResource
from fdai.core.wara.report import SUPPLEMENTAL_SOURCES, build_wara_report
from fdai.delivery.wara_report import WaraFileReportSink, read_collector_context
from fdai.rule_catalog.schema.framework_catalog import load_framework_catalog
from fdai.rule_catalog.schema.wara_assessment import load_wara_assessment_catalog
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[4]
SUBSCRIPTION = "00000000-0000-0000-0000-000000000000"


@pytest.fixture
def assets():
    root = ROOT / "rule-catalog/collected/wara-aprl"
    framework = load_framework_catalog(root, best_practices=(), objective_refs=frozenset())[0]
    catalog, _ = load_wara_assessment_catalog(
        root / "assessment/crosswalk.json",
        root / "assessment/queries.json",
        framework=framework,
        framework_path=root / "azure-wara.json",
    )
    at = datetime(2026, 10, 30, tzinfo=UTC)
    request = WaraAssessmentRequest(
        assessment_id="wara-assessment:report-test",
        framework_revision=catalog.source_revision,
        crosswalk_digest=catalog.crosswalk_digest,
        ontology_release="ontology-example",
        inventory_generation="inventory-example",
        workload_id="workload-example",
        resources=(
            WaraScopedResource(
                f"/subscriptions/{SUBSCRIPTION}/resourceGroups/rg-example"
                "/providers/Microsoft.Compute/virtualMachines/vm-example",
                "Microsoft.Compute/virtualMachines",
            ),
        ),
        evaluated_at=at,
        recorded_at=at,
    )
    return catalog, request, WaraAssessmentRuntime(catalog).assess(request)


def _collector():
    raw = {"scriptDetails": {"ImplicitSubscriptionIds": [f"/subscriptions/{SUBSCRIPTION}"]}}
    for index, field in enumerate(
        ("advisory", "outages", "retirements", "supportTickets", "serviceHealth")
    ):
        raw[field] = [{"detail": "private-context-not-for-export"} for _ in range(index)]
    return raw


def test_missing_supplemental_sources_stay_explicitly_unavailable(assets) -> None:
    catalog, _, result = assets
    report = build_wara_report(result, catalog)
    rows = report.widgets[2].data["rows"]
    assert {row["source"] for row in rows} == set(SUPPLEMENTAL_SOURCES)
    assert all(row["status"] == "unavailable" and row["count"] is None for row in rows)


@pytest.mark.parametrize("change", ["authority", "digest", "omission", "revision"])
def test_export_rejects_authority_or_incomplete_tampered_assessments(assets, change) -> None:
    catalog, _, result = assets
    changes = {
        "authority": {"execution_authority": True},
        "digest": {"result_digest": "sha256:" + "0" * 64},
        "omission": {"controls": result.controls[:-1]},
        "revision": {"framework_revision": "0" * 40},
    }
    with pytest.raises(ValueError, match="WARA report"):
        build_wara_report(replace(result, **changes[change]), catalog)


def test_collector_context_never_claims_verified_completeness(assets, tmp_path) -> None:
    _, request, _ = assets
    path = tmp_path / "collector.json"
    path.write_text(json.dumps(_collector()))
    rows = read_collector_context(path, request)
    assert len(rows) == 5
    assert [row["count"] for row in rows] == list(range(5))
    assert all(row["status"] == "unverified" and row["complete"] is False for row in rows)
    assert "private-context-not-for-export" not in json.dumps(rows)


@pytest.mark.parametrize("change", ["wider_scope", "missing_scope", "wrong_shape"])
def test_collector_context_cannot_widen_scope_or_hide_invalid_data(
    assets, tmp_path, change
) -> None:
    _, request, _ = assets
    raw = _collector()
    if change == "wider_scope":
        raw["scriptDetails"]["ImplicitSubscriptionIds"].append(
            "/subscriptions/00000000-0000-0000-0000-000000000001"
        )
    elif change == "missing_scope":
        raw.pop("scriptDetails")
    else:
        raw["supportTickets"] = "invalid"
    path = tmp_path / "collector.json"
    path.write_text(json.dumps(raw))
    with pytest.raises(ValueError, match="WARA collector context"):
        read_collector_context(path, request)


async def test_job_report_sink_exports_all_393_controls_in_all_formats(assets, tmp_path) -> None:
    catalog, request, result = assets
    context = tmp_path / "collector.json"
    context.write_text(json.dumps(_collector()))
    directory = tmp_path / "reports"
    await WaraFileReportSink(output_directory=directory, collector_context=context).write(
        result, catalog, request
    )
    paths = {path.suffix: path for path in directory.iterdir()}
    assert set(paths) == {".json", ".csv", ".xlsx", ".pptx"}
    assert all(stat.S_IMODE(path.stat().st_mode) == 0o600 for path in paths.values())
    report = json.loads(paths[".json"].read_text())
    assert len(report["widgets"][1]["data"]["rows"]) == 393
    assert report["variables"]["execution_authority"] == "false"
    assert "private-context-not-for-export" not in paths[".json"].read_text()
    with paths[".csv"].open() as stream:
        rows = list(csv.DictReader(stream))
    assert len([row for row in rows if row["widget_id"] == "recommendations"]) == 393
    book = load_workbook(paths[".xlsx"])
    assert len([row for row in book["Report"].values if row[0] == "recommendations"]) == 393
    slides = Presentation(paths[".pptx"])
    text = "\n".join(
        shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert all(control.recommendation_id in text for control in result.controls)
    assert "unknown" in text
    assert result.result_digest in text
